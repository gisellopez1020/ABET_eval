"""
Lectura de la primera hoja de un .xlsx con openpyxl, con las mismas reglas que tenía
pandas.read_excel (sheet_name=0) en las importaciones:

- La primera hoja del libro, no la activa.
- Fórmulas con su valor calculado guardado en el archivo.
- Celdas vacías y celdas fusionadas (salvo la primera) como None.
- Números enteros guardados como float (2021001.0) como int: un código no gana ".0".
- Sin las filas ni columnas vacías del final (Excel guarda celdas con formato y sin valor);
  las filas vacías intermedias se conservan para que "Fila N" coincida con la de Excel.
- Todas las filas con el mismo ancho: el de la fila más larga, encabezado incluido.
"""
import io
from typing import Any, List

from openpyxl import load_workbook


class ExcelIlegible(Exception):
    """El contenido no es un .xlsx que openpyxl pueda abrir."""


def _valor(v: Any) -> Any:
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def leer_primera_hoja(contenido: bytes) -> List[List[Any]]:
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelIlegible(str(exc) or type(exc).__name__) from exc
    try:
        filas = [[_valor(v) for v in fila] for fila in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()

    while filas and all(v is None for v in filas[-1]):
        filas.pop()
    ancho = max(
        (max(i for i, v in enumerate(fila) if v is not None) + 1 for fila in filas if any(v is not None for v in fila)),
        default=0,
    )
    return [(fila + [None] * ancho)[:ancho] for fila in filas]
