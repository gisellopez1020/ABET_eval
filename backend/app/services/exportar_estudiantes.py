"""
Libro Excel con la lista de estudiantes (openpyxl, sin base de datos).

Columnas: Nombre | Apellido(s) | Número de ID | Dirección de correo | Grupo
[| Asignatura, solo si las filas vienen de más de una asignatura].

El nombre se guarda completo en la base (nombre_completo); aquí se separa solo para
el archivo con separar_nombre(). El correo sale tal cual del campo email: si es
NULL la celda queda vacía, nunca se deduce del nombre o del código.
"""
import io
import re
from dataclasses import dataclass
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ENCABEZADOS = ["Nombre", "Apellido(s)", "Número de ID", "Dirección de correo", "Grupo"]
COLUMNA_ASIGNATURA = "Asignatura"
ANCHOS = [24, 28, 16, 34, 12, 36]

# Partículas de apellidos compuestos (De La Cruz, Del Valle, San Martín): se pegan
# a la palabra siguiente para no partir el apellido
PARTICULAS = {"de", "del", "la", "las", "los", "y", "san"}

_BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
_FONDO_ENCABEZADO = PatternFill("solid", fgColor="E7E6E6")


@dataclass
class FilaEstudiante:
    nombre_completo: str
    codigo: str
    email: Optional[str]
    grupo: str
    asignatura: str


def separar_nombre(nombre_completo: str) -> tuple[str, str]:
    """
    (nombre, apellidos): los dos últimos bloques son los apellidos y el resto el nombre.
    Con dos bloques, uno y uno; con uno, todo es nombre. Un bloque es una palabra más
    las partículas que la preceden ("De La Cruz" es un solo bloque).

    Limitación conocida: dos nombres y un solo apellido ("Juan Pablo Escobar") quedan
    como "Juan" / "Pablo Escobar"; con solo el texto no se puede distinguir de
    "Juan Escobar Duarte".
    """
    bloques: list[list[str]] = []
    for palabra in nombre_completo.split():
        if bloques and bloques[-1][-1].lower() in PARTICULAS:
            bloques[-1].append(palabra)
        else:
            bloques.append([palabra])

    if len(bloques) <= 1:
        return " ".join(nombre_completo.split()), ""
    n_apellidos = 2 if len(bloques) >= 3 else 1
    nombre = " ".join(p for b in bloques[:-n_apellidos] for p in b)
    apellidos = " ".join(p for b in bloques[-n_apellidos:] for p in b)
    return nombre, apellidos


def libro_estudiantes(filas: Sequence[FilaEstudiante]) -> bytes:
    """Una hoja "Estudiantes" con una fila por estudiante, en el orden recibido."""
    con_asignatura = len({f.asignatura for f in filas}) > 1
    encabezados = ENCABEZADOS + ([COLUMNA_ASIGNATURA] if con_asignatura else [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Estudiantes"
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = _FONDO_ENCABEZADO
        celda.border = _BORDE
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for f in filas:
        nombre, apellidos = separar_nombre(f.nombre_completo)
        valores = [nombre, apellidos, f.codigo, f.email or None, f.grupo]
        if con_asignatura:
            valores.append(f.asignatura)
        ws.append(valores)

    for i, ancho in enumerate(ANCHOS[: len(encabezados)]):
        ws.column_dimensions[chr(ord("A") + i)].width = ancho
    # Número de ID como texto: un código como 0012 no debe perder los ceros
    for celda in ws["C"][1:]:
        celda.number_format = "@"
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def nombre_archivo_estudiantes(partes: Sequence[str]) -> str:
    """Estudiantes[_{codigo curso}][_{sección}].xlsx sin caracteres inválidos para archivos."""
    base = "_".join(["Estudiantes", *partes])
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f\s]+', "_", base)
    return re.sub(r"_+", "_", base).strip("_") + ".xlsx"
