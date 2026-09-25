"""
Libros Excel de Estadísticas ABET por actividad (openpyxl, sin base de datos).

- Hoja "Conteo": una fila por Criterio ABET con la cantidad de estudiantes en
  cada rango del curso y una torta por fila.
- Hojas de detalle (una por equipo o estudiante): mismo formato que la rúbrica
  del profesor (Punto del informe / Aspecto evaluado / Criterio / Cumple /
  Nota informe / Nota del proyecto / Cantidad estudiantes). Los valores son
  fijos, no fórmulas: la app es la fuente de verdad.
"""
import io
import re
from dataclasses import dataclass
from itertools import groupby
from typing import Hashable, Optional, Sequence

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.reporte import RangoReporte, ReporteCriterioItem

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HOJA_CONTEO = "Conteo"
SIN_CLASIFICAR = "Sin clasificar"
COLOR_SIN_CLASIFICAR = "9CA3AF"
# Misma escala que colorRango del frontend: rojo (rango más bajo) -> ámbar -> verde
ESCALA_RANGOS = ["C8102E", "FFB300", "2E7D32"]

ENCABEZADOS_DETALLE = [
    "Punto del informe", "Aspecto evaluado", "Criterio", "Cumple",
    "Nota informe", "Nota del proyecto", "Cantidad estudiantes",
]
ANCHOS_DETALLE = {"A": 14, "B": 24, "C": 80, "D": 10, "E": 12, "F": 12, "G": 13}
FILA_DATOS = 3  # los encabezados ocupan las filas 1-2 (combinadas)

_BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
_FONDO_ENCABEZADO = PatternFill("solid", fgColor="E7E6E6")
_CENTRADO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_IZQUIERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)


@dataclass
class AspectoDetalle:
    clave_bloque: Hashable            # aspectos seguidos con la misma clave comparten Nota informe
    codigo_abet: Optional[str]
    nombre: str
    criterios: list[tuple[str, Optional[int]]]   # (texto, cumple 0/1 o None si no se calificó)
    nota: Optional[float]             # nota del bloque (None si el bloque está incompleto)


@dataclass
class HojaDetalle:
    nombre: str
    aspectos: list[AspectoDetalle]
    nota_proyecto: Optional[float]    # None si la calificación está incompleta
    cantidad_estudiantes: int
    integrantes: list[str]


def color_rango(indice: int, total: int) -> str:
    """Espejo de colorRango (frontend/src/utils/rangos.ts)."""
    if total <= 1:
        return ESCALA_RANGOS[-1]
    t = indice / (total - 1) * (len(ESCALA_RANGOS) - 1)
    tramo = min(int(t), len(ESCALA_RANGOS) - 2)
    f = t - tramo
    a, b = ESCALA_RANGOS[tramo], ESCALA_RANGOS[tramo + 1]
    canales = (
        round(int(a[i:i + 2], 16) + (int(b[i:i + 2], 16) - int(a[i:i + 2], 16)) * f) for i in (0, 2, 4)
    )
    return "".join(f"{c:02X}" for c in canales)


def nombre_archivo(curso, actividad, sufijo: str) -> str:
    """ABET_{codigo}_{actividad}_{periodo}_{sufijo}.xlsx sin caracteres inválidos para archivos."""
    base = f"ABET_{curso.codigo}_{actividad.nombre}_{curso.periodo}_{sufijo}"
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f\s]+', "_", base)
    return re.sub(r"_+", "_", base).strip("_") + ".xlsx"


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    """Máximo 31 caracteres, sin []:*?/\\ y único (sin distinguir mayúsculas)."""
    base = re.sub(r"[\[\]:*?/\\]", "", nombre).strip().strip("'")[:31] or "Hoja"
    candidato, n = base, 2
    while candidato.lower() in usados:
        sufijo = f" ({n})"
        candidato = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(candidato.lower())
    return candidato


def _estilo(ws: Worksheet, rango: str, alineacion: Alignment, borde: bool = True) -> None:
    for fila in ws[rango]:
        for celda in fila:
            celda.alignment = alineacion
            if borde:
                celda.border = _BORDE


def _escribir_conteo(ws: Worksheet, rangos: Sequence[RangoReporte], criterios: Sequence[ReporteCriterioItem]) -> None:
    con_sin_clasificar = any(c.sin_clasificar for c in criterios)
    columnas = [r.etiqueta for r in rangos] + ([SIN_CLASIFICAR] if con_sin_clasificar else [])
    colores = [color_rango(i, len(rangos)) for i in range(len(rangos))] + (
        [COLOR_SIN_CLASIFICAR] if con_sin_clasificar else []
    )
    ultima_col = len(columnas) + 1

    ws.append(["Calificación", *columnas])
    for c in criterios:
        ws.append([f"ABET {c.codigo}", *(c.rangos.get(r.etiqueta, 0) for r in rangos),
                   *([c.sin_clasificar] if con_sin_clasificar else [])])

    ws.column_dimensions["A"].width = 16
    for i in range(2, ultima_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = _FONDO_ENCABEZADO
    _estilo(ws, f"A1:{get_column_letter(ultima_col)}{len(criterios) + 1}", _CENTRADO)

    # Una torta por fila, apiladas a la derecha de la tabla
    ancla = get_column_letter(ultima_col + 2)
    categorias = Reference(ws, min_col=2, max_col=ultima_col, min_row=1, max_row=1)
    for i, c in enumerate(criterios):
        fila = i + 2
        torta = PieChart()
        torta.title = f"ABET {c.codigo}"
        torta.add_data(Reference(ws, min_col=2, max_col=ultima_col, min_row=fila, max_row=fila), from_rows=True)
        torta.set_categories(categorias)
        torta.dataLabels = DataLabelList()
        torta.dataLabels.showPercent = True
        torta.dataLabels.numFmt = "0%;;;"   # oculta las etiquetas de porciones vacías (0 %)
        for j, color in enumerate(colores):
            punto = DataPoint(idx=j)
            punto.graphicalProperties.solidFill = color
            torta.series[0].dPt.append(punto)
        torta.width, torta.height = 12, 7.5
        ws.add_chart(torta, f"{ancla}{1 + i * 16}")


def _combinar(ws: Worksheet, columna: str, inicio: int, fin: int, valor) -> None:
    ws[f"{columna}{inicio}"] = valor
    if fin > inicio:
        ws.merge_cells(f"{columna}{inicio}:{columna}{fin}")


def _escribir_detalle(ws: Worksheet, hoja: HojaDetalle) -> None:
    for i, encabezado in enumerate(ENCABEZADOS_DETALLE, start=1):
        col = get_column_letter(i)
        _combinar(ws, col, 1, 2, encabezado)
        ws[f"{col}1"].font = Font(bold=True)
        for fila in (1, 2):
            ws[f"{col}{fila}"].fill = _FONDO_ENCABEZADO
    for col, ancho in ANCHOS_DETALLE.items():
        ws.column_dimensions[col].width = ancho

    # Filas de criterios y rango de filas de cada aspecto
    fila = FILA_DATOS
    tramos: list[tuple[AspectoDetalle, int, int]] = []
    for aspecto in hoja.aspectos:
        inicio = fila
        for texto, cumple in aspecto.criterios:
            ws[f"C{fila}"] = texto
            ws[f"D{fila}"] = cumple
            fila += 1
        tramos.append((aspecto, inicio, fila - 1))
    ultima = max(fila - 1, FILA_DATOS)

    # B: un aspecto; A: aspectos seguidos con el mismo código; E: aspectos seguidos del mismo bloque
    for aspecto, inicio, fin in tramos:
        _combinar(ws, "B", inicio, fin, aspecto.nombre)
    for codigo, grupo in groupby(tramos, key=lambda t: t[0].codigo_abet):
        grupo = list(grupo)
        _combinar(ws, "A", grupo[0][1], grupo[-1][2], f"ABET {codigo}" if codigo else None)
    for _, grupo in groupby(tramos, key=lambda t: t[0].clave_bloque):
        grupo = list(grupo)
        _combinar(ws, "E", grupo[0][1], grupo[-1][2], grupo[0][0].nota)

    _combinar(ws, "F", FILA_DATOS, ultima, hoja.nota_proyecto)
    ws[f"G{FILA_DATOS}"] = hoja.cantidad_estudiantes

    _estilo(ws, "A1:G2", _CENTRADO)
    _estilo(ws, f"A{FILA_DATOS}:G{ultima}", _CENTRADO)
    _estilo(ws, f"B{FILA_DATOS}:C{ultima}", _IZQUIERDA)
    for col in ("E", "F"):
        for (celda,) in ws[f"{col}{FILA_DATOS}:{col}{ultima}"]:
            celda.number_format = "0.00"

    # Integrantes debajo de la tabla, dejando una fila libre
    for i, integrante in enumerate(hoja.integrantes):
        ws[f"C{ultima + 2 + i}"] = integrante


def _guardar(wb: Workbook) -> bytes:
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def libro_resumen(rangos: Sequence[RangoReporte], criterios: Sequence[ReporteCriterioItem]) -> bytes:
    wb = Workbook()
    wb.active.title = HOJA_CONTEO
    _escribir_conteo(wb.active, rangos, criterios)
    return _guardar(wb)


def libro_detalle(
    rangos: Sequence[RangoReporte],
    criterios: Sequence[ReporteCriterioItem],
    hojas: Sequence[HojaDetalle],
) -> bytes:
    """Hoja Conteo (igual que el resumen) seguida de una hoja por equipo o estudiante."""
    wb = Workbook()
    wb.active.title = HOJA_CONTEO
    _escribir_conteo(wb.active, rangos, criterios)
    usados = {HOJA_CONTEO.lower()}
    for hoja in hojas:
        _escribir_detalle(wb.create_sheet(_nombre_hoja(hoja.nombre, usados)), hoja)
    return _guardar(wb)
