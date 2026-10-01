"""
Exportación de estudiantes a Excel (GET /estudiantes/exportar-excel), separación
nombre/apellidos solo al exportar (services/exportar_estudiantes.separar_nombre),
correo real opcional en la importación y la migración 0006 (estudiantes.email).

El correo nunca se calcula: sale del campo email o la celda queda vacía.
Reutiliza las fixtures de SQLite en memoria.
"""
import io
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.models import Curso, Estudiante, Seccion
from app.services.exportar_estudiantes import separar_nombre
from tests.test_catalogo_ra_abet import MOCK_USER, client, db_session  # noqa: F401 (fixtures)

BACKEND = Path(__file__).resolve().parent.parent
ENCABEZADOS = ["Nombre", "Apellido(s)", "Número de ID", "Dirección de correo", "Grupo"]
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── separar_nombre ───────────────────────────────────────────────────────────

# Los 20 nombres reales del seed (scripts/seed_545210.json) con la división esperada
SEED = [
    ("Valentina Rojas Medina", "Valentina", "Rojas Medina"),
    ("Juan Esteban Ospina Cárdenas", "Juan Esteban", "Ospina Cárdenas"),
    ("Camila Andrea Rincón Silva", "Camila Andrea", "Rincón Silva"),
    ("Santiago Moreno Vidal", "Santiago", "Moreno Vidal"),
    ("Daniela Torres Cifuentes", "Daniela", "Torres Cifuentes"),
    ("Mateo Alejandro Ruiz Gómez", "Mateo Alejandro", "Ruiz Gómez"),
    ("Isabella Castañeda Ríos", "Isabella", "Castañeda Ríos"),
    ("Nicolás Herrera Quintero", "Nicolás", "Herrera Quintero"),
    ("Sara Valentina Gómez Peña", "Sara Valentina", "Gómez Peña"),
    ("Juan Pablo Escobar Duarte", "Juan Pablo", "Escobar Duarte"),
    ("Laura Sofía Muñoz Aguirre", "Laura Sofía", "Muñoz Aguirre"),
    ("David Felipe Cárdenas Ospina", "David Felipe", "Cárdenas Ospina"),
    ("María José Zapata Correa", "María José", "Zapata Correa"),
    ("Andrés Camilo Pardo Vélez", "Andrés Camilo", "Pardo Vélez"),
    ("Paula Andrea Salazar Vanegas", "Paula Andrea", "Salazar Vanegas"),
    ("Kevin Stiven Loaiza Ramírez", "Kevin Stiven", "Loaiza Ramírez"),
    ("Mariana Restrepo Duque", "Mariana", "Restrepo Duque"),
    ("Sebastián Naranjo Vega", "Sebastián", "Naranjo Vega"),
    ("Luisa Fernanda Bonilla Trejos", "Luisa Fernanda", "Bonilla Trejos"),
    ("Cristian David Marulanda Ibarra", "Cristian David", "Marulanda Ibarra"),
]


def test_la_tabla_cubre_exactamente_los_nombres_del_seed():
    datos = json.loads((BACKEND / "scripts" / "seed_545210.json").read_text(encoding="utf-8"))
    del_seed = [n for e in datos["equipos"] for n in e["estudiantes"]]
    assert sorted(del_seed) == sorted(n for n, _, _ in SEED)


@pytest.mark.parametrize("completo,nombre,apellidos", SEED, ids=[s[0] for s in SEED])
def test_nombres_del_seed(completo, nombre, apellidos):
    assert separar_nombre(completo) == (nombre, apellidos)
    # La importación los guarda en mayúsculas: la división no depende del caso
    assert separar_nombre(completo.upper()) == (nombre.upper(), apellidos.upper())


@pytest.mark.parametrize("completo,esperado", [
    ("María José De La Cruz Pérez", ("María José", "De La Cruz Pérez")),
    ("ANA DEL VALLE ROJAS", ("ANA", "DEL VALLE ROJAS")),
    ("Luis San Martín Díaz", ("Luis", "San Martín Díaz")),
    ("María De Los Ángeles Pérez Gómez", ("María De Los Ángeles", "Pérez Gómez")),
    ("Pedro Gómez", ("Pedro", "Gómez")),
    ("Cher", ("Cher", "")),
    ("  Valentina   Rojas  Medina ", ("Valentina", "Rojas Medina")),
    ("", ("", "")),
])
def test_casos_especiales(completo, esperado):
    assert separar_nombre(completo) == esperado


def test_limitacion_conocida_dos_nombres_un_apellido():
    """Documentado en separar_nombre: sin más datos no se puede distinguir."""
    assert separar_nombre("Juan Pablo Escobar") == ("Juan", "Pablo Escobar")


# ── Exportar ─────────────────────────────────────────────────────────────────

def _curso(db_session, nombre="Redes", codigo="R-1", email=MOCK_USER["email"], secciones=("S1",)):
    curso = Curso(nombre=nombre, codigo=codigo, periodo="2026-2", docente_email=email, ra_abet=[])
    db_session.add(curso)
    db_session.flush()
    secs = [Seccion(nombre=s, curso_id=curso.id) for s in secciones]
    db_session.add_all(secs)
    db_session.commit()
    return curso, secs


def _est(db_session, seccion, nombre, codigo, email=None):
    est = Estudiante(nombre_completo=nombre, codigo_estudiante=codigo, email=email, seccion_id=seccion.id)
    db_session.add(est)
    db_session.commit()
    return est


def _exportar(client, **params):
    resp = client.get("/estudiantes/exportar-excel", params=params)
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == MIME_XLSX
    ws = load_workbook(io.BytesIO(resp.content)).active
    return resp, [list(r) for r in ws.iter_rows(values_only=True)]


class TestExportar:
    def test_cinco_columnas_en_orden_y_correo_vacio_si_no_hay(self, client, db_session):
        _, [s1] = _curso(db_session)
        _est(db_session, s1, "VALENTINA ROJAS MEDINA", "2021001", "valentina.rojas@uao.edu.co")
        _est(db_session, s1, "JUAN ESTEBAN OSPINA CÁRDENAS", "2021002")

        _, filas = _exportar(client)
        assert filas[0] == ENCABEZADOS
        assert filas[1:] == [
            ["JUAN ESTEBAN", "OSPINA CÁRDENAS", "2021002", None, "S1"],
            ["VALENTINA", "ROJAS MEDINA", "2021001", "valentina.rojas@uao.edu.co", "S1"],
        ]

    def test_correo_vacio_no_se_inventa_a_partir_del_codigo(self, client, db_session):
        _, [s1] = _curso(db_session)
        _est(db_session, s1, "ANA RUIZ", "2021002")
        _, filas = _exportar(client)
        assert filas[1][3] is None
        assert not any("@" in str(c) for c in filas[1] if c)

    def test_grupo_es_el_nombre_real_de_la_seccion_y_orden(self, client, db_session):
        _, [s1, s2] = _curso(db_session, secciones=("Grupo 01", "Grupo 02"))
        _est(db_session, s2, "BETO DIAZ", "2")
        _est(db_session, s1, "ZOE ARIAS", "3")
        _est(db_session, s1, "ANA RUIZ", "1")
        _, filas = _exportar(client)
        assert [(f[0], f[4]) for f in filas[1:]] == [("ANA", "Grupo 01"), ("ZOE", "Grupo 01"), ("BETO", "Grupo 02")]

    def test_codigo_se_exporta_como_texto(self, client, db_session):
        _, [s1] = _curso(db_session)
        _est(db_session, s1, "ANA RUIZ", "0012")
        _, filas = _exportar(client)
        assert filas[1][2] == "0012"

    def test_columna_asignatura_solo_con_varias_asignaturas(self, client, db_session):
        redes, [r1] = _curso(db_session, "Redes", "R-1")
        _, [b1] = _curso(db_session, "Bases de Datos", "BD-1")
        _est(db_session, r1, "ANA RUIZ", "1")
        _est(db_session, b1, "BETO DIAZ", "2")

        _, filas = _exportar(client)
        assert filas[0] == ENCABEZADOS + ["Asignatura"]
        assert [f[5] for f in filas[1:]] == ["Bases de Datos (BD-1 · 2026-2)", "Redes (R-1 · 2026-2)"]

        # Filtrado a una asignatura: vuelve a las 5 columnas
        _, filas = _exportar(client, curso_id=redes.id)
        assert filas[0] == ENCABEZADOS
        assert [f[0] for f in filas[1:]] == ["ANA"]

    def test_filtro_por_seccion_y_nombre_del_archivo(self, client, db_session):
        _, [s1, s2] = _curso(db_session, secciones=("S1", "S2"))
        _est(db_session, s1, "ANA RUIZ", "1")
        _est(db_session, s2, "BETO DIAZ", "2")
        resp, filas = _exportar(client, seccion_id=s2.id)
        assert [f[0] for f in filas[1:]] == ["BETO"]
        assert 'filename="Estudiantes_R-1_S2.xlsx"' in resp.headers["content-disposition"]

    def test_solo_estudiantes_del_docente(self, client, db_session):
        _, [mia] = _curso(db_session)
        _, [ajena] = _curso(db_session, "Ajena", "A-1", email="otro@uao.edu.co")
        _est(db_session, mia, "ANA RUIZ", "1")
        _est(db_session, ajena, "ZOE ARIAS", "9")
        _, filas = _exportar(client)
        assert [f[0] for f in filas[1:]] == ["ANA"]
        assert len(filas[0]) == 5  # la asignatura ajena no cuenta para la columna Asignatura

    def test_seccion_archivada_solo_si_se_pide(self, client, db_session):
        _, [activa, archivada] = _curso(db_session, secciones=("S1", "Vieja"))
        archivada.activo = False
        db_session.commit()
        _est(db_session, activa, "ANA RUIZ", "1")
        _est(db_session, archivada, "BETO DIAZ", "2")
        _, filas = _exportar(client)
        assert [f[0] for f in filas[1:]] == ["ANA"]
        _, filas = _exportar(client, seccion_id=archivada.id)
        assert [f[0] for f in filas[1:]] == ["BETO"]

    def test_sin_estudiantes_solo_encabezados(self, client, db_session):
        _curso(db_session)
        _, filas = _exportar(client)
        assert filas == [ENCABEZADOS]


# ── Importar con correo ──────────────────────────────────────────────────────

def _xlsx(filas):
    wb = Workbook()
    for fila in filas:
        wb.active.append(list(fila))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _correos(db_session, seccion):
    db_session.expire_all()
    return {
        e.codigo_estudiante: e.email
        for e in db_session.query(Estudiante).filter_by(seccion_id=seccion.id).all()
    }


class TestImportarCorreo:
    @pytest.mark.parametrize("encabezado", ["Email", "Correo", "E-mail", "Correo electrónico"])
    def test_csv_con_columna_de_correo(self, client, db_session, encabezado):
        _, [s1] = _curso(db_session)
        csv = f"Nombre,Codigo,{encabezado}\nAna Ruiz,1, Ana.Ruiz@UAO.edu.co \nBeto Diaz,2,\n"
        resp = client.post(f"/secciones/{s1.id}/estudiantes/csv",
                           files={"archivo": ("e.csv", csv.encode(), "text/csv")})
        assert resp.json() == {"importados": 2, "errores": []}
        assert _correos(db_session, s1) == {"1": "ana.ruiz@uao.edu.co", "2": None}

    def test_csv_sin_columna_de_correo_queda_null(self, client, db_session):
        _, [s1] = _curso(db_session)
        resp = client.post(f"/secciones/{s1.id}/estudiantes/csv",
                           files={"archivo": ("e.csv", b"Nombre,Codigo\nAna Ruiz,1\n", "text/csv")})
        assert resp.json()["importados"] == 1
        assert _correos(db_session, s1) == {"1": None}

    def test_excel_con_columna_de_correo(self, client, db_session):
        _, [s1] = _curso(db_session)
        contenido = _xlsx([("Nombre", "Código", "Correo"), ("Ana Ruiz", 1, "ana@uao.edu.co"), ("Beto Diaz", 2, None)])
        resp = client.post(f"/secciones/{s1.id}/estudiantes/csv", files={"archivo": ("e.xlsx", contenido, MIME_XLSX)})
        assert resp.json() == {"importados": 2, "errores": []}
        assert _correos(db_session, s1) == {"1": "ana@uao.edu.co", "2": None}

    def test_correo_invalido_importa_al_estudiante_sin_correo_y_avisa(self, client, db_session):
        _, [s1] = _curso(db_session)
        csv = "Nombre,Codigo,Email\nAna Ruiz,1,ana ruiz\nBeto Diaz,2,beto@uao.edu.co\n"
        resp = client.post(f"/secciones/{s1.id}/estudiantes/csv",
                           files={"archivo": ("e.csv", csv.encode(), "text/csv")})
        assert resp.json() == {
            "importados": 2,
            "errores": ["Fila 2: correo 'ana ruiz' no válido, se deja en blanco"],
        }
        assert _correos(db_session, s1) == {"1": None, "2": "beto@uao.edu.co"}

    def test_vista_previa_excel_incluye_el_correo(self, client, db_session):
        _, [s1] = _curso(db_session)
        contenido = _xlsx([("Nombre", "Codigo", "Email"), ("Ana Ruiz", "1", "ana@uao.edu.co")])
        resp = client.post(f"/secciones/{s1.id}/estudiantes/vista-previa",
                           files={"archivo": ("e.xlsx", contenido, MIME_XLSX)})
        assert resp.json()["estudiantes"] == [{"nombre": "ANA RUIZ", "codigo": "1", "email": "ana@uao.edu.co"}]

    def test_importado_con_correo_sale_en_el_export(self, client, db_session):
        _, [s1] = _curso(db_session)
        csv = "Nombre,Codigo,Email\nValentina Rojas Medina,1,v@uao.edu.co\n"
        client.post(f"/secciones/{s1.id}/estudiantes/csv", files={"archivo": ("e.csv", csv.encode(), "text/csv")})
        _, filas = _exportar(client, seccion_id=s1.id)
        assert filas[1] == ["VALENTINA", "ROJAS MEDINA", "1", "v@uao.edu.co", "S1"]


# ── Alta manual con correo ───────────────────────────────────────────────────

def _crear(client, seccion, **extra):
    return client.post(
        f"/secciones/{seccion.id}/estudiantes",
        json={"nombre_completo": "Ana Ruiz", "codigo_estudiante": "1", **extra},
    )


class TestAltaManualCorreo:
    def test_con_correo_lo_guarda_y_lo_devuelve_en_el_listado(self, client, db_session):
        _, [s1] = _curso(db_session)
        resp = _crear(client, s1, email="  Ana.Ruiz@UAO.edu.co ")
        assert resp.status_code == 201
        assert (resp.json()["email"], resp.json()["aviso"]) == ("ana.ruiz@uao.edu.co", None)

        [listado] = client.get(f"/secciones/{s1.id}/estudiantes").json()
        assert (listado["nombre_completo"], listado["email"]) == ("ANA RUIZ", "ana.ruiz@uao.edu.co")
        assert _correos(db_session, s1) == {"1": "ana.ruiz@uao.edu.co"}

    @pytest.mark.parametrize("extra", [{}, {"email": None}, {"email": ""}, {"email": "   "}],
                             ids=["omitido", "null", "vacio", "espacios"])
    def test_sin_correo_sigue_funcionando_con_email_none(self, client, db_session, extra):
        _, [s1] = _curso(db_session)
        resp = _crear(client, s1, **extra)
        assert resp.status_code == 201
        datos = resp.json()
        assert (datos["nombre_completo"], datos["codigo_estudiante"], datos["seccion_id"]) == ("ANA RUIZ", "1", s1.id)
        assert (datos["email"], datos["aviso"]) == (None, None)
        assert client.get(f"/secciones/{s1.id}/estudiantes").json()[0]["email"] is None

    @pytest.mark.parametrize("invalido", ["ana ruiz", "ana@", "ana.uao.edu.co", "a@b"])
    def test_correo_invalido_no_impide_la_creacion(self, client, db_session, invalido):
        _, [s1] = _curso(db_session)
        resp = _crear(client, s1, email=invalido)
        assert resp.status_code == 201
        assert resp.json()["email"] is None
        assert resp.json()["aviso"] == f"El correo '{invalido}' no es válido, se dejó en blanco"
        # El estudiante existe, con el correo en blanco: ni rechazado ni inventado
        assert _correos(db_session, s1) == {"1": None}

    def test_creado_a_mano_con_correo_sale_en_el_export(self, client, db_session):
        _, [s1] = _curso(db_session)
        _crear(client, s1, email="ana@uao.edu.co")
        _, filas = _exportar(client, seccion_id=s1.id)
        assert filas[1] == ["ANA", "RUIZ", "1", "ana@uao.edu.co", "S1"]


# ── Migración 0006 ───────────────────────────────────────────────────────────

def _alembic_sql(*args):
    """
    SQL que generaría Alembic en modo offline (--sql, sin base de datos). Se usa el
    ejecutable instalado y no `import alembic`: la carpeta backend/alembic/ (con
    __init__.py) tapa al paquete cuando los tests corren desde backend/.
    """
    exe = Path(sys.executable).parent / ("alembic.exe" if os.name == "nt" else "alembic")
    resultado = subprocess.run(
        [str(exe), *args, "--sql"], cwd=BACKEND, capture_output=True, text=True, timeout=60,
    )
    assert resultado.returncode == 0, resultado.stderr
    return resultado.stdout


def test_migracion_0006_agrega_email_nullable_sin_default():
    sql = _alembic_sql("upgrade", "0005:0006")
    assert "ALTER TABLE estudiantes ADD COLUMN email VARCHAR(254);" in sql
    assert "NOT NULL" not in sql and "DEFAULT" not in sql.upper()
    assert "version_num='0006'" in sql


def test_migracion_0006_downgrade_quita_email():
    sql = _alembic_sql("downgrade", "0006:0005")
    assert "ALTER TABLE estudiantes DROP COLUMN email;" in sql


def test_0006_es_la_ultima_revision():
    assert "0006 (head)" in subprocess.run(
        [str(Path(sys.executable).parent / ("alembic.exe" if os.name == "nt" else "alembic")), "heads"],
        cwd=BACKEND, capture_output=True, text=True, timeout=60,
    ).stdout
