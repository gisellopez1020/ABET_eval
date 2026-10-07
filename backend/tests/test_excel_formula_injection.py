"""
Inyección de fórmulas en los Excel exportados (CWE-1236).

openpyxl guarda como fórmula cualquier texto que empiece con "=": un nombre de
estudiante o un texto de rúbrica importado como =HYPERLINK(...) quedaba vivo en el
libro. El texto importado se escribe como texto literal (data_type "s" y el prefijo
oculto quotePrefix), sin cambiar el valor: la lista exportada se puede reimportar.
"+", "-" y "@" ya los guarda openpyxl como texto; se prueban para que siga así.
"""
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.schemas.reporte import RangoReporte, ReporteCriterioItem
from app.services.exportar_estudiantes import FilaEstudiante, libro_estudiantes
from app.services.lista_estudiantes import leer_lista
from app.services.reporte_excel import AspectoDetalle, HojaDetalle, libro_detalle
from app.utils.excel_seguro import escribir_texto, texto_seguro_para_excel

FORMULA = '=HYPERLINK("HTTP://EVIL.EXAMPLE","CLIC")'
PELIGROSOS = [FORMULA, "+1+1", "-2+3", "@SUM(A1)"]
NORMAL = "Presenta el informe completo"


def _abrir(contenido: bytes):
    return load_workbook(io.BytesIO(contenido))


def _es_texto_literal(celda, valor):
    assert celda.value == valor
    assert celda.data_type == "s"


# ── Helper ───────────────────────────────────────────────────────────────────

class TestHelper:
    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_texto_peligroso_queda_como_texto(self, valor):
        ws = Workbook().active
        escribir_texto(ws, "A1", valor)
        _es_texto_literal(ws["A1"], valor)

    def test_formula_lleva_prefijo_de_texto(self):
        ws = Workbook().active
        escribir_texto(ws, "A1", FORMULA)
        assert ws["A1"].quotePrefix is True

    def test_texto_normal_no_se_toca(self):
        ws = Workbook().active
        escribir_texto(ws, "A1", NORMAL)
        _es_texto_literal(ws["A1"], NORMAL)
        assert ws["A1"].quotePrefix is False

    @pytest.mark.parametrize("valor,tipo", [(3.5, "n"), (1, "n"), (True, "b"), (None, "n")])
    def test_valores_no_texto_no_se_tocan(self, valor, tipo):
        ws = Workbook().active
        ws["A1"] = valor
        texto_seguro_para_excel(ws["A1"])
        assert ws["A1"].value == valor and ws["A1"].data_type == tipo
        assert ws["A1"].quotePrefix is False


# ── Lista de estudiantes ─────────────────────────────────────────────────────

def _libro_estudiantes(**campos):
    base = dict(nombre_completo="ANA ROJAS MEDINA", codigo="001", email=None, grupo="G1", asignatura="A")
    filas = [FilaEstudiante(**{**base, **campos}), FilaEstudiante(**{**base, "asignatura": "B"})]
    return _abrir(libro_estudiantes(filas)).active


class TestListaEstudiantes:
    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_nombre(self, valor):
        ws = _libro_estudiantes(nombre_completo=f"{valor} ROJAS MEDINA")
        _es_texto_literal(ws["A2"], valor)
        _es_texto_literal(ws["B2"], "ROJAS MEDINA")

    def test_apellidos(self):
        ws = _libro_estudiantes(nombre_completo=f"ANA {FORMULA} MEDINA")
        _es_texto_literal(ws["B2"], f"{FORMULA} MEDINA")

    @pytest.mark.parametrize("columna,campo", [("C", "codigo"), ("D", "email"), ("E", "grupo"), ("F", "asignatura")])
    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_demas_columnas(self, columna, campo, valor):
        ws = _libro_estudiantes(**{campo: valor})
        _es_texto_literal(ws[f"{columna}2"], valor)

    def test_texto_normal_no_se_toca(self):
        ws = _libro_estudiantes()
        assert [c.value for c in ws[2]] == ["ANA", "ROJAS MEDINA", "001", None, "G1", "A"]
        assert not any(c.quotePrefix for c in ws[2])

    def test_reimportar_conserva_el_nombre(self):
        nombre = f"{FORMULA} ROJAS MEDINA"
        contenido = libro_estudiantes([FilaEstudiante(nombre, "001", None, "G1", "A")])
        validos, _ = leer_lista(contenido, "lista.xlsx", None)
        assert validos == [(nombre, "001", None)]


# ── Reporte de estadísticas ──────────────────────────────────────────────────

def _libro_detalle(texto=NORMAL, aspecto=NORMAL, integrante="ANA (001)", etiqueta="0.0-5.0"):
    rangos = [RangoReporte(etiqueta=etiqueta, minimo=0.0, maximo=5.0)]
    criterios = [ReporteCriterioItem(
        codigo="1.1", descripcion="d", codigo_padre=None, peso=None,
        rangos={etiqueta: 1}, sin_clasificar=0, total=1,
    )]
    hoja = HojaDetalle(
        nombre="Equipo 1",
        aspectos=[AspectoDetalle(
            clave_bloque="1.1", codigo_abet="1.1", nombre=aspecto,
            criterios=[(texto, 1), ("otro criterio", 0)], nota=2.5,
        )],
        nota_proyecto=2.5, cantidad_estudiantes=1, integrantes=[integrante],
    )
    wb = _abrir(libro_detalle(rangos, criterios, [hoja]))
    return wb["Conteo"], wb["Equipo 1"]


class TestReporteDetalle:
    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_texto_criterio(self, valor):
        _, ws = _libro_detalle(texto=valor)
        _es_texto_literal(ws["C3"], valor)

    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_nombre_aspecto(self, valor):
        _, ws = _libro_detalle(aspecto=valor)
        _es_texto_literal(ws["B3"], valor)

    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_integrante(self, valor):
        _, ws = _libro_detalle(integrante=valor)
        _es_texto_literal(ws["C6"], valor)

    @pytest.mark.parametrize("valor", PELIGROSOS)
    def test_etiqueta_rango(self, valor):
        conteo, _ = _libro_detalle(etiqueta=valor)
        _es_texto_literal(conteo["B1"], valor)

    def test_texto_normal_y_numeros_no_se_tocan(self):
        conteo, ws = _libro_detalle()
        assert [conteo["A1"].value, conteo["B1"].value, conteo["B2"].value] == ["Calificación", "0.0-5.0", 1]
        assert [ws["A3"].value, ws["B3"].value, ws["C3"].value, ws["C6"].value] == [
            "ABET 1.1", NORMAL, NORMAL, "ANA (001)",
        ]
        assert (ws["D3"].data_type, ws["E3"].data_type, ws["F3"].data_type, ws["G3"].data_type) == ("n",) * 4
        assert ws["E3"].number_format == "0.00"
        assert not any(c.quotePrefix for fila in ws.iter_rows() for c in fila)
