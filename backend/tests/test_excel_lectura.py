"""
Lectura de la primera hoja (utils/excel_lectura.py), que reemplaza a pandas.read_excel en las
importaciones. Cada caso usa un .xlsx real generado con openpyxl; los valores esperados son los
que devolvía pandas 2.2.2 con el mismo archivo (salvo los tipos: pandas los infería por columna).
"""
import datetime
import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from app.utils.excel_lectura import ExcelIlegible, leer_primera_hoja


def _xlsx(filas, preparar=None):
    wb = Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(list(fila))
    if preparar:
        preparar(wb, ws)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


class TestLeerPrimeraHoja:
    def test_valores_y_tipos(self):
        filas = leer_primera_hoja(_xlsx([
            ("Codigo", "Nota", "Activo", "Fecha"),
            ("0012", 3.5, True, datetime.datetime(2024, 1, 2)),
            (2021001, 2021002.0, None, None),
        ]))
        assert filas == [
            ["Codigo", "Nota", "Activo", "Fecha"],
            ["0012", 3.5, True, datetime.datetime(2024, 1, 2)],
            [2021001, 2021002, None, None],   # el float entero llega como int: sin ".0"
        ]
        assert type(filas[2][1]) is int

    def test_celdas_fusionadas_llegan_como_none(self):
        filas = leer_primera_hoja(_xlsx(
            [("Aspecto", "Criterio"), ("Diseño", "A"), (None, "B")],
            preparar=lambda wb, ws: ws.merge_cells("A2:A3"),
        ))
        assert filas[1:] == [["Diseño", "A"], [None, "B"]]

    def test_conserva_filas_vacias_intermedias_y_quita_las_del_final(self):
        def formato_al_final(wb, ws):
            # Celdas con formato y sin valor: Excel las guarda, pandas las ignoraba
            for fila in range(ws.max_row + 1, ws.max_row + 4):
                ws.cell(fila, 1).font = Font(bold=True)
            ws.cell(1, 9).font = Font(bold=True)

        filas = leer_primera_hoja(_xlsx(
            [("Nombre", "Codigo"), ("Ana", "1"), (None, None), ("Eva", "3")], preparar=formato_al_final,
        ))
        assert filas == [["Nombre", "Codigo"], ["Ana", "1"], [None, None], ["Eva", "3"]]

    def test_ancho_de_la_fila_mas_larga(self):
        filas = leer_primera_hoja(_xlsx([("A", "C"), ("X", "c1", 100)]))
        assert filas == [["A", "C", None], ["X", "c1", 100]]

    def test_lee_la_primera_hoja_aunque_no_sea_la_activa(self):
        def otra_hoja_primero(wb, ws):
            primera = wb.create_sheet("Primera", 0)
            primera.append(["Nombre", "Codigo"])
            primera.append(["DE LA PRIMERA", "9"])
            wb.active = 1

        filas = leer_primera_hoja(_xlsx([("Nombre", "Codigo"), ("DE LA ACTIVA", "1")], preparar=otra_hoja_primero))
        assert filas[1] == ["DE LA PRIMERA", "9"]

    def test_formula_sin_valor_guardado_llega_vacia(self):
        # openpyxl no calcula: usa el valor que guardó Excel (este archivo no tiene ninguno)
        assert leer_primera_hoja(_xlsx([("Nombre", "Codigo"), ("Ana", "=1+1")]))[1] == ["Ana", None]

    def test_libro_vacio(self):
        assert leer_primera_hoja(_xlsx([])) == []

    def test_archivo_que_no_es_excel(self):
        with pytest.raises(ExcelIlegible):
            leer_primera_hoja(b"Nombre,Codigo\nAna,1\n")
