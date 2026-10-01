"""
EstudianteService (services/estudiantes.py) probado sin HTTP: se llama al servicio con
una sesión SQLite y se verifica el error de negocio que lanza, o el resultado con su aviso.
"""
import ast
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.models import Estudiante
from app.services.errores import (
    ArchivoInvalido, AsignaturaNoEncontrada, EstudianteNoEncontrado, SeccionFueraDeAsignatura,
    SeccionNoEncontrada, SinPermiso,
)
from app.services.estudiantes import EstudianteConAviso, EstudianteService, ResultadoImportacion
from tests.test_catalogo_ra_abet import MOCK_USER, db_session  # noqa: F401 (fixtures)
from tests.test_exportar_estudiantes import _curso, _est

EMAIL = MOCK_USER["email"]
OTRO = "otro@uao.edu.co"
CSV = "text/csv"


@pytest.fixture()
def servicio(db_session):
    return EstudianteService(db_session)


@pytest.fixture()
def s1(db_session):
    _, [seccion] = _curso(db_session)
    return seccion


@pytest.fixture()
def ajena(db_session):
    _, [seccion] = _curso(db_session, "Ajena", "A-1", email=OTRO)
    return seccion


def _guardados(db_session, seccion):
    db_session.expire_all()
    return sorted(
        (e.nombre_completo, e.codigo_estudiante, e.email)
        for e in db_session.query(Estudiante).filter_by(seccion_id=seccion.id)
    )


# ── Sección inexistente o ajena: todas las operaciones por sección ───────────

OPERACIONES_POR_SECCION = {
    "listar": lambda s, sid: s.listar(sid, EMAIL),
    "agregar": lambda s, sid: s.agregar(sid, EMAIL, "Ana", "1", None),
    "importar": lambda s, sid: s.importar(sid, EMAIL, b"Nombre,Codigo\nAna,1\n", "l.csv", CSV),
    "vista_previa": lambda s, sid: s.vista_previa(sid, EMAIL, b"Nombre,Codigo\nAna,1\n", "l.csv", CSV),
}


@pytest.mark.parametrize("operacion", OPERACIONES_POR_SECCION)
def test_seccion_inexistente(servicio, operacion):
    with pytest.raises(SeccionNoEncontrada):
        OPERACIONES_POR_SECCION[operacion](servicio, 99999)


@pytest.mark.parametrize("operacion", OPERACIONES_POR_SECCION)
def test_seccion_de_otro_docente_no_guarda_nada(servicio, db_session, ajena, operacion):
    with pytest.raises(SinPermiso, match="^No tiene permiso sobre esta sección$"):
        OPERACIONES_POR_SECCION[operacion](servicio, ajena.id)
    assert _guardados(db_session, ajena) == []


# ── Listar, agregar y la regla del correo ────────────────────────────────────

def test_listar_con_promedio(servicio, db_session, s1):
    _est(db_session, s1, "ANA", "1")
    [ana] = servicio.listar(s1.id, EMAIL)
    assert (ana.nombre_completo, ana.promedio) == ("ANA", None)


class TestAgregar:
    def test_normaliza_y_no_avisa(self, servicio, db_session, s1):
        r = servicio.agregar(s1.id, EMAIL, "  Ana Ruiz ", " 1 ", "  Ana@UAO.edu.co ")
        assert isinstance(r, EstudianteConAviso)
        assert r.aviso is None
        assert _guardados(db_session, s1) == [("ANA RUIZ", "1", "ana@uao.edu.co")]

    def test_correo_invalido_crea_igual_y_avisa(self, servicio, db_session, s1):
        r = servicio.agregar(s1.id, EMAIL, "Ana", "1", "ana@")
        assert r.estudiante.email is None
        assert r.aviso == "El correo 'ana@' no es válido, se dejó en blanco"
        assert _guardados(db_session, s1) == [("ANA", "1", None)]


class TestEditar:
    def test_solo_cambia_lo_enviado(self, servicio, db_session, s1):
        ana = _est(db_session, s1, "ANA", "1", "ana@uao.edu.co")
        r = servicio.editar(ana.id, EMAIL, {"nombre_completo": "Ana Ruiz"})
        assert r.aviso is None
        assert _guardados(db_session, s1) == [("ANA RUIZ", "1", "ana@uao.edu.co")]

    def test_correo_none_lo_borra(self, servicio, db_session, s1):
        ana = _est(db_session, s1, "ANA", "1", "ana@uao.edu.co")
        assert servicio.editar(ana.id, EMAIL, {"email": None}).estudiante.email is None

    def test_correo_invalido_conserva_el_anterior(self, servicio, db_session, s1):
        ana = _est(db_session, s1, "ANA", "1", "ana@uao.edu.co")
        r = servicio.editar(ana.id, EMAIL, {"email": "ana uao"})
        assert r.estudiante.email == "ana@uao.edu.co"
        assert r.aviso == "El correo 'ana uao' no es válido, se conservó el anterior"

    def test_correo_invalido_sin_anterior_queda_en_blanco(self, servicio, db_session, s1):
        ana = _est(db_session, s1, "ANA", "1")
        r = servicio.editar(ana.id, EMAIL, {"email": "ana uao"})
        assert r.estudiante.email is None
        assert r.aviso == "El correo 'ana uao' no es válido, se dejó en blanco"

    def test_inexistente(self, servicio):
        with pytest.raises(EstudianteNoEncontrado, match="^Estudiante no encontrado$"):
            servicio.editar(99999, EMAIL, {"nombre_completo": "X"})

    def test_de_otro_docente_no_cambia(self, servicio, db_session, ajena):
        zoe = _est(db_session, ajena, "ZOE", "9")
        with pytest.raises(SinPermiso, match="^No tiene permiso para editar este estudiante$"):
            servicio.editar(zoe.id, EMAIL, {"nombre_completo": "X"})
        assert _guardados(db_session, ajena) == [("ZOE", "9", None)]


class TestEliminar:
    def test_elimina(self, servicio, db_session, s1):
        ana = _est(db_session, s1, "ANA", "1")
        servicio.eliminar(ana.id, EMAIL)
        assert _guardados(db_session, s1) == []

    def test_inexistente(self, servicio):
        with pytest.raises(EstudianteNoEncontrado):
            servicio.eliminar(99999, EMAIL)

    def test_de_otro_docente_no_borra(self, servicio, db_session, ajena):
        zoe = _est(db_session, ajena, "ZOE", "9")
        with pytest.raises(SinPermiso, match="^No tiene permiso para eliminar este estudiante$"):
            servicio.eliminar(zoe.id, EMAIL)
        assert _guardados(db_session, ajena) == [("ZOE", "9", None)]


# ── Importar y vista previa ──────────────────────────────────────────────────

class TestImportar:
    def test_guarda_los_validos_y_devuelve_los_avisos(self, servicio, db_session, s1):
        csv = b"Nombre,Codigo,Grupo\nAna,1,S1\nBeto,2,S9\n,3,S1\n"
        r = servicio.importar(s1.id, EMAIL, csv, "lista.csv", CSV)
        assert r == ResultadoImportacion(
            importados=2,
            avisos=[
                "Fila 4: nombre o código vacío, se omite",
                "1 fila tiene Grupo 'S9', distinto de esta sección ('S1'); se importa igual",
            ],
        )
        assert _guardados(db_session, s1) == [("ANA", "1", None), ("BETO", "2", None)]

    def test_archivo_invalido_no_guarda_nada(self, servicio, db_session, s1):
        with pytest.raises(ArchivoInvalido):
            servicio.importar(s1.id, EMAIL, b"Nombre,Correo\nAna,a@x.co\n", "lista.csv", CSV)
        assert _guardados(db_session, s1) == []

    def test_vista_previa_no_guarda(self, servicio, db_session, s1):
        vista = servicio.vista_previa(s1.id, EMAIL, b"Nombre,Codigo\nAna,1\n", "lista.csv", CSV)
        assert [e.nombre for e in vista.estudiantes] == ["ANA"]
        assert _guardados(db_session, s1) == []


# ── Exportar ─────────────────────────────────────────────────────────────────

class TestExportar:
    def test_libro_y_nombre_de_archivo(self, servicio, db_session):
        curso, [s1] = _curso(db_session)
        _est(db_session, s1, "ANA RUIZ", "1")
        contenido, nombre = servicio.exportar(EMAIL, curso_id=curso.id, seccion_id=s1.id)
        assert nombre == "Estudiantes_R-1_S1.xlsx"
        filas = list(load_workbook(io.BytesIO(contenido)).active.iter_rows(values_only=True))
        assert filas[1][:3] == ("ANA", "RUIZ", "1")

    def test_asignatura_inexistente(self, servicio):
        with pytest.raises(AsignaturaNoEncontrada, match="^Asignatura no encontrada$"):
            servicio.exportar(EMAIL, curso_id=99999)

    def test_asignatura_de_otro_docente(self, servicio, ajena):
        with pytest.raises(SinPermiso, match="^No tiene permiso sobre esta asignatura$"):
            servicio.exportar(EMAIL, curso_id=ajena.curso_id)

    def test_seccion_de_otra_asignatura(self, servicio, db_session):
        redes, _ = _curso(db_session, "Redes", "R-1")
        _, [b1] = _curso(db_session, "Bases de Datos", "BD-1")
        with pytest.raises(SeccionFueraDeAsignatura, match="^Sección no encontrada en esta asignatura$"):
            servicio.exportar(EMAIL, curso_id=redes.id, seccion_id=b1.id)


# ── Independencia de FastAPI ─────────────────────────────────────────────────

@pytest.mark.parametrize("modulo", ["app/services/estudiantes.py", "app/services/lista_estudiantes.py"])
def test_no_depende_de_fastapi(modulo):
    """El servicio y la lectura de listas no importan nada de FastAPI ni de Starlette."""
    arbol = ast.parse((Path(__file__).parent.parent / modulo).read_text(encoding="utf-8"))
    importados = set()
    for nodo in ast.walk(arbol):
        if isinstance(nodo, ast.Import):
            importados.update(alias.name for alias in nodo.names)
        elif isinstance(nodo, ast.ImportFrom) and nodo.module:
            importados.add(nodo.module)
    raices = {nombre.split(".")[0] for nombre in importados}
    assert not raices & {"fastapi", "starlette"}
