"""
Estadísticas ABET por actividad: reporte acotado a una actividad, libro resumen
(hoja Conteo con tortas), libro de detalle con el formato de la rúbrica del
profesor y sincronización con Google Drive (simulada, fallida y exitosa).

Reutiliza las fixtures y utilidades de test_reportes (SQLite en memoria).
"""
import base64
import io
import re
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest

from app.config import settings
from app.models import EquipoTrabajo, MiembroEquipo
from app.models.actividad import TipoActividad
from app.services import google_drive
from app.services.reporte_excel import _nombre_hoja, color_rango, nombre_archivo
from tests.test_catalogo_ra_abet import MOCK_USER, client, db_session  # noqa: F401 (fixtures)
from tests.test_reportes import DEFAULT, _actividad, _calificar, _estudiante, _seccion, curso  # noqa: F401

RANGOS_DEFAULT = ["0.0-2.9", "3.0-3.9", "4.0-5.0"]


def _url(curso, actividad, sufijo=""):
    return f"/reportes/abet/{curso.id}/actividad/{actividad.id}{sufijo}"


def _equipo(db_session, curso, actividad, nombre, miembros):
    equipo = EquipoTrabajo(nombre=nombre, actividad_id=actividad.id, seccion_id=_seccion(db_session, curso).id)
    db_session.add(equipo)
    db_session.flush()
    db_session.add_all([MiembroEquipo(equipo_id=equipo.id, estudiante_id=m.id) for m in miembros])
    db_session.commit()
    return equipo


def _libro(contenido: bytes):
    return openpyxl.load_workbook(io.BytesIO(contenido))


def _cantidad_graficos(contenido: bytes) -> int:
    """openpyxl no lee gráficos al abrir un libro: se cuentan las partes del zip."""
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        return sum(1 for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))


def _detalle(client, curso, actividad, **body):
    resp = client.post(_url(curso, actividad, "/detalle-xlsx"), json=body)
    assert resp.status_code == 200, resp.text
    datos = resp.json()
    return datos, base64.b64decode(datos["archivo_base64"])


@pytest.fixture()
def grupal(db_session, curso):
    """
    Actividad grupal: dos aspectos 2.1.1 seguidos (un bloque), uno sin vincular y uno 4.1.1.
    Equipo "Equipo: 1/A" (ANA, BEA) calificado completo; equipo sin calificar.
    """
    act, crits = _actividad(
        db_session, curso,
        [("2.1.1", [20, 20]), ("2.1.1", [20]), (None, [20]), ("4.1.1", [20])],
        tipo=TipoActividad.grupal,
    )
    ana, bea = _estudiante(db_session, curso, "ANA"), _estudiante(db_session, curso, "BEA")
    equipo = _equipo(db_session, curso, act, "Equipo: 1/A", [ana, bea])
    _equipo(db_session, curso, act, "Sin calificar", [_estudiante(db_session, curso, "CAR")])
    # 2.1.1: 5*(1*20+0*20+1*20)/60 = 3.33 | sin vincular: 5.0 | 4.1.1: 0.0
    _calificar(db_session, crits[0] + crits[1] + crits[2] + crits[3], [1, 0, 1, 1, 0], equipo=equipo)
    return act


class TestReporteActividad:
    def test_acotado_a_la_actividad(self, client, db_session, curso):
        act1, c1 = _actividad(db_session, curso, [("2.1.1", [100])])
        act2, c2 = _actividad(db_session, curso, [("2.1.1", [50]), ("2.1.2", [50])])
        ana = _estudiante(db_session, curso, "ANA")
        _calificar(db_session, c1[0], [1], estudiante=ana)            # 5.0 en act1
        _calificar(db_session, c2[0] + c2[1], [0, 1], estudiante=ana)  # 0.0 y 5.0 en act2

        rep = client.get(_url(curso, act1)).json()
        assert rep["actividad_id"] == act1.id
        assert rep["actividad_tipo"] == "individual"
        assert [c["codigo"] for c in rep["criterios"]] == ["2.1.1"]
        assert rep["criterios"][0]["rangos"] == {**DEFAULT, "4.0-5.0": 1}
        # En el RA solo cuenta lo vinculado en esta actividad
        assert rep["resultados"][0]["criterios_con_evidencia"] == ["2.1.1"]
        assert rep["resultados"][0]["criterios_sin_evidencia"] == ["2.1.2", "2.1.3"]

        rep2 = client.get(_url(curso, act2)).json()
        assert rep2["criterios"][0]["rangos"] == {**DEFAULT, "0.0-2.9": 1}

    def test_grupal_cuenta_por_integrante(self, client, db_session, curso, grupal):
        rep = client.get(_url(curso, grupal)).json()
        c = {i["codigo"]: i for i in rep["criterios"]}
        assert c["2.1.1"]["rangos"] == {**DEFAULT, "3.0-3.9": 2}
        assert c["4.1.1"]["rangos"] == {**DEFAULT, "0.0-2.9": 2}

    def test_filtro_por_seccion(self, client, db_session, curso):
        act, crits = _actividad(db_session, curso, [("2.1.1", [100])])
        _calificar(db_session, crits[0], [1], estudiante=_estudiante(db_session, curso, "ANA", "S1"))
        _calificar(db_session, crits[0], [0], estudiante=_estudiante(db_session, curso, "BEA", "S2"))
        s2 = _seccion(db_session, curso, "S2")
        rep = client.get(_url(curso, act), params={"seccion_id": s2.id}).json()
        assert rep["criterios"][0]["rangos"] == {**DEFAULT, "0.0-2.9": 1}

    def test_actividad_de_otro_curso(self, client, db_session, curso):
        from app.models import Curso
        otro = Curso(nombre="X", codigo="X", periodo="2026-2", docente_email=MOCK_USER["email"], ra_abet=[])
        db_session.add(otro)
        db_session.commit()
        act, _ = _actividad(db_session, otro, [("2.1.1", [100])])
        for sufijo in ("", "/resumen-xlsx"):
            assert client.get(_url(curso, act, sufijo)).status_code == 404
        assert client.post(_url(curso, act, "/detalle-xlsx"), json={}).status_code == 404

    def test_seccion_de_otro_curso(self, client, db_session, curso):
        act, _ = _actividad(db_session, curso, [("2.1.1", [100])])
        assert client.get(_url(curso, act), params={"seccion_id": 9999}).status_code == 404

    def test_curso_de_otro_docente(self, client, db_session, curso):
        from app.models import Curso
        otro = Curso(nombre="X", codigo="X", periodo="2026-2", docente_email="otro@uao.edu.co", ra_abet=[])
        db_session.add(otro)
        db_session.commit()
        act, _ = _actividad(db_session, otro, [("2.1.1", [100])])
        assert client.get(_url(otro, act)).status_code == 403


class TestResumenXlsx:
    def test_hoja_conteo_con_una_torta_por_criterio(self, client, db_session, curso, grupal):
        resp = client.get(_url(curso, grupal, "/resumen-xlsx"))
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert "filename*=UTF-8''ABET_R-1_Act_2026-2_resumen.xlsx" in resp.headers["content-disposition"]

        wb = _libro(resp.content)
        assert wb.sheetnames == ["Conteo"]
        filas = [list(f) for f in wb["Conteo"].iter_rows(values_only=True)]
        assert filas == [
            ["Calificación", *RANGOS_DEFAULT],
            ["ABET 2.1.1", 0, 2, 0],
            ["ABET 4.1.1", 2, 0, 0],
        ]
        assert _cantidad_graficos(resp.content) == 2

    def test_usa_los_rangos_del_curso(self, client, db_session, curso, grupal):
        curso.rangos_calificacion = [
            {"etiqueta": "Bajo", "minimo": 0, "maximo": 2.9},
            {"etiqueta": "Alto", "minimo": 3, "maximo": 5},
        ]
        db_session.commit()
        wb = _libro(client.get(_url(curso, grupal, "/resumen-xlsx")).content)
        assert [c.value for c in wb["Conteo"][1]] == ["Calificación", "Bajo", "Alto"]


class TestDetalleXlsx:
    def test_hoja_por_equipo_con_formato_de_la_rubrica(self, client, db_session, curso, grupal):
        datos, contenido = _detalle(client, curso, grupal)
        assert datos["nombre_archivo"] == "ABET_R-1_Act_2026-2_detalle.xlsx"

        wb = _libro(contenido)
        # El equipo sin calificaciones no lleva hoja; el nombre se limpia de caracteres inválidos
        assert wb.sheetnames == ["Conteo", "Equipo 1A"]
        assert _cantidad_graficos(contenido) == 2
        ws = wb["Equipo 1A"]

        assert [ws.cell(1, c).value for c in range(1, 8)] == [
            "Punto del informe", "Aspecto evaluado", "Criterio", "Cumple",
            "Nota informe", "Nota del proyecto", "Cantidad estudiantes",
        ]
        combinadas = {str(r) for r in ws.merged_cells.ranges}
        assert {"A1:A2", "G1:G2"} <= combinadas
        # Dos aspectos 2.1.1 seguidos: A y E combinadas sobre ambos; B por aspecto
        assert {"A3:A5", "B3:B4", "E3:E5", "F3:F7"} <= combinadas
        assert ws["A3"].value == "ABET 2.1.1"
        assert ws["A6"].value is None and ws["A7"].value == "ABET 4.1.1"
        assert [ws[f"D{f}"].value for f in range(3, 8)] == [1, 0, 1, 1, 0]
        assert ws["E3"].value == pytest.approx(5 * 40 / 60)
        assert ws["E6"].value == pytest.approx(5.0)   # aspecto sin vincular: su propia nota
        assert ws["E7"].value == pytest.approx(0.0)
        assert ws["F3"].value == pytest.approx(3.0)   # suma de nota_calculada: 3 de 5 criterios al 20 %
        assert ws["G3"].value == 2
        assert [ws["C9"].value, ws["C10"].value] == ["ANA (ANA)", "BEA (BEA)"]

    def test_calificacion_incompleta_deja_vacios(self, client, db_session, curso):
        act, crits = _actividad(db_session, curso, [("2.1.1", [50, 50])])
        _calificar(db_session, crits[0][:1], [1], estudiante=_estudiante(db_session, curso, "ANA"))
        _estudiante(db_session, curso, "BEA")   # sin calificaciones: sin hoja
        _, contenido = _detalle(client, curso, act)
        wb = _libro(contenido)
        assert wb.sheetnames == ["Conteo", "ANA"]
        ws = wb["ANA"]
        assert (ws["D3"].value, ws["D4"].value) == (1, None)
        assert ws["E3"].value is None and ws["F3"].value is None
        assert ws["G3"].value == 1


class TestNombreConSeccion:
    def test_resumen_con_y_sin_seccion(self, client, db_session, curso, grupal):
        s1 = _seccion(db_session, curso, "S1")
        sin = client.get(_url(curso, grupal, "/resumen-xlsx")).headers["content-disposition"]
        con = client.get(_url(curso, grupal, "/resumen-xlsx"), params={"seccion_id": s1.id}).headers["content-disposition"]
        assert "filename*=UTF-8''ABET_R-1_Act_2026-2_resumen.xlsx" in sin
        assert "filename*=UTF-8''ABET_R-1_Act_S1_2026-2_resumen.xlsx" in con

    def test_detalle_con_y_sin_seccion(self, client, db_session, curso, grupal):
        s1 = _seccion(db_session, curso, "S1")
        assert _detalle(client, curso, grupal)[0]["nombre_archivo"] == "ABET_R-1_Act_2026-2_detalle.xlsx"
        con, _ = _detalle(client, curso, grupal, seccion_id=s1.id)
        assert con["nombre_archivo"] == "ABET_R-1_Act_S1_2026-2_detalle.xlsx"


class TestSincronizacionDrive:
    def test_simulada_con_skip_auth(self, client, db_session, curso, grupal, monkeypatch):
        monkeypatch.setattr(settings, "skip_auth", True)
        llamadas = []
        monkeypatch.setattr(google_drive, "_subir_archivo", lambda *a: llamadas.append(a))
        datos, contenido = _detalle(client, curso, grupal)
        assert datos["drive"] == {"estado": "simulado", "detalle": None, "enlace": None}
        assert llamadas == []
        assert _libro(contenido).sheetnames[0] == "Conteo"

    def test_sin_token_devuelve_el_archivo_y_marca_error(self, client, db_session, curso, grupal, monkeypatch):
        monkeypatch.setattr(settings, "skip_auth", False)
        monkeypatch.setattr(google_drive, "obtener_tokens_drive", lambda email: None)
        datos, contenido = _detalle(client, curso, grupal)
        assert datos["drive"]["estado"] == "error"
        assert "iniciar sesión" in datos["drive"]["detalle"]
        assert "Equipo 1A" in _libro(contenido).sheetnames

    def test_fallo_de_subida_no_bloquea(self, client, db_session, curso, grupal, monkeypatch):
        monkeypatch.setattr(settings, "skip_auth", False)
        monkeypatch.setattr(google_drive, "obtener_tokens_drive", lambda email: {"access_token": "t"})

        def _falla(*_):
            raise ConnectionError("sin conexión")

        monkeypatch.setattr(google_drive, "_subir_archivo", _falla)
        datos, contenido = _detalle(client, curso, grupal)
        assert datos["drive"]["estado"] == "error"
        assert "sin conexión" in datos["drive"]["detalle"]
        assert len(contenido) > 0

    def test_subida_exitosa(self, client, db_session, curso, grupal, monkeypatch):
        monkeypatch.setattr(settings, "skip_auth", False)
        monkeypatch.setattr(google_drive, "obtener_tokens_drive", lambda email: {"access_token": "t"})
        recibido = {}

        def _sube(token, email, nombre, contenido, mimetype):
            recibido.update(email=email, nombre=nombre, mimetype=mimetype, tamano=len(contenido))
            return {"id": "abc", "webViewLink": "https://drive.google.com/file/d/abc"}

        monkeypatch.setattr(google_drive, "_subir_archivo", _sube)
        datos, _ = _detalle(client, curso, grupal)
        assert datos["drive"] == {
            "estado": "sincronizado", "detalle": None, "enlace": "https://drive.google.com/file/d/abc",
        }
        assert recibido["email"] == MOCK_USER["email"]
        assert recibido["nombre"] == "ABET_R-1_Act_2026-2_detalle.xlsx"
        assert recibido["mimetype"].startswith("application/vnd.openxmlformats")
        assert recibido["tamano"] > 0

    def test_escapa_comillas_en_la_busqueda(self, monkeypatch):
        consultas = []

        class _Files:
            def list(self, q, fields):
                consultas.append(q)
                return SimpleNamespace(execute=lambda: {"files": []})

            def create(self, **_):
                return SimpleNamespace(execute=lambda: {"id": "x", "webViewLink": "l"})

        monkeypatch.setattr(google_drive, "_build_drive_service", lambda t: SimpleNamespace(files=_Files))
        monkeypatch.setattr(google_drive, "_obtener_o_crear_carpeta", lambda s, e: "carpeta")
        google_drive._subir_archivo("t", "e", "Proyecto d'Alembert.xlsx", b"x", "m")
        assert "name = 'Proyecto d\\'Alembert.xlsx'" in consultas[0]


class TestUtilidades:
    def test_nombre_hoja_unico_y_acotado(self):
        usados = {"conteo"}
        assert _nombre_hoja("Conteo", usados) == "Conteo (2)"
        largo = "Nombre de estudiante muy largo que excede"
        primero = _nombre_hoja(largo, usados)
        segundo = _nombre_hoja(largo, usados)
        assert len(primero) == 31 and len(segundo) == 31
        assert segundo.endswith("(2)") and primero != segundo

    def test_nombre_archivo_sin_caracteres_invalidos(self):
        curso = SimpleNamespace(codigo="R-1", periodo="2026-2")
        actividad = SimpleNamespace(nombre='Proyecto: fase 1/2 "final"')
        assert nombre_archivo(curso, actividad, "detalle") == "ABET_R-1_Proyecto_fase_1_2_final_2026-2_detalle.xlsx"

    def test_nombre_archivo_incluye_la_seccion(self):
        curso = SimpleNamespace(codigo="R-1", periodo="2026-2")
        actividad = SimpleNamespace(nombre="Proyecto final")
        seccion = SimpleNamespace(nombre="Grupo 1")
        assert nombre_archivo(curso, actividad, "resumen", seccion) == "ABET_R-1_Proyecto_final_Grupo_1_2026-2_resumen.xlsx"
        assert nombre_archivo(curso, actividad, "resumen", None) == "ABET_R-1_Proyecto_final_2026-2_resumen.xlsx"

    def test_color_rango_igual_al_frontend(self):
        assert [color_rango(i, 3) for i in range(3)] == ["C8102E", "FFB300", "2E7D32"]
        assert color_rango(0, 1) == "2E7D32"
